from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import get_session
from app.main import create_app
from app.models import OrderLine, ShipmentLine, ShipmentReport, User
from app.services.orders import create_order_line
from app.services.packing_drafts import create_packing_draft, submit_packing_draft
from app.services.skus import upsert_sku_mapping
from app.services.shipments import submit_shipment_report

PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"a" * 256


def _client(db_session):
    admin = User(username="api_more_admin", display_name="老板", password_hash="x", role="admin", is_active=True)
    worker = User(username="api_more_worker", display_name="仓库A", password_hash="x", role="worker", is_active=True)
    db_session.add_all([admin, worker])
    db_session.commit()
    app = create_app()
    app.dependency_overrides[get_session] = lambda: db_session
    client = TestClient(app)
    client.cookies.set("zy_user_id", str(admin.id))
    return client, admin, worker


def test_dashboard_returns_stats(db_session):
    client, _admin, worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)
    submit_shipment_report(
        db_session,
        worker.id,
        "2026-08-04",
        "源兴发",
        "裁判",
        "圆领裁判",
        [{"size": "M", "quantity": 10}],
        [],
        "",
    )

    response = client.get("/api/v1/dashboard")

    assert response.status_code == 200
    payload = response.json()
    assert payload["pending"] == 1
    assert len(payload["recent"]) >= 1


def test_orders_options_and_create_order(db_session):
    client, _admin, _worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 50)

    options = client.get("/api/v1/orders/options")
    assert options.status_code == 200
    assert "源兴发" in options.json()["companies"]
    assert options.json()["choices"]["源兴发"]["裁判"]["圆领裁判"] == ["M"]

    created = client.post(
        "/api/v1/orders",
        data={
            "company_name": "源兴发",
            "product_name": "裁判",
            "style_name": "圆领裁判",
            "sizes": ["M", "L"],
            "quantities": ["10", "20"],
            "order_date": "2026-08-05",
            "delivery_date": "",
            "accessories": "帽子",
            "material": "涤纶",
            "spec_size": "110CM",
            "note": "",
        },
    )
    assert created.status_code == 200
    payload = created.json()
    assert payload["created"] == 2
    assert payload["duplicates"] == []

    duplicate = client.post(
        "/api/v1/orders",
        data={
            "company_name": "源兴发",
            "product_name": "裁判",
            "style_name": "圆领裁判",
            "sizes": ["M"],
            "quantities": ["10"],
            "order_date": "2026-08-05",
            "delivery_date": "",
        },
    )
    assert duplicate.json()["duplicates"]


def test_review_pending_approve_reject(db_session):
    client, _admin, worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)
    report = submit_shipment_report(
        db_session,
        worker.id,
        "2026-08-04",
        "源兴发",
        "裁判",
        "圆领裁判",
        [{"size": "M", "quantity": 100}],
        [],
        "",
    )
    assert report.status == "pending_review"

    pending = client.get("/api/v1/review/pending")
    assert pending.status_code == 200
    assert len(pending.json()["reports"]) == 1

    response = client.post(f"/api/v1/review/{report.id}/approve", data={"note": "ok"})
    assert response.status_code == 200
    db_session.refresh(report)
    assert report.status == "approved_after_edit"


def test_shipments_list_filter_and_waybill_update(db_session):
    client, _admin, worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)
    report = submit_shipment_report(
        db_session,
        worker.id,
        "2026-08-04",
        "源兴发",
        "裁判",
        "圆领裁判",
        [{"size": "M", "quantity": 10}],
        [],
        "",
    )
    report.waybill_no = "YT111"
    db_session.commit()

    response = client.get("/api/v1/shipments?waybill=YT111")
    assert response.status_code == 200
    payload = response.json()
    assert payload["total"] == 1
    assert payload["reports"][0]["waybill_no"] == "YT111"

    response = client.post(f"/api/v1/shipments/{report.id}/waybill", data={"waybill_no": "YT999"})
    assert response.status_code == 200
    db_session.refresh(report)
    assert report.waybill_no == "YT999"


def test_shipments_list_includes_bound_order_reference_per_line(db_session):
    client, _admin, worker = _client(db_session)
    dated_order = create_order_line(
        db_session,
        "源兴发",
        "裁判",
        "圆领裁判",
        "M",
        100,
        order_date="2026-07-15",
    )
    batched_order = create_order_line(
        db_session,
        "源兴发",
        "裁判",
        "圆领裁判",
        "L",
        100,
        order_date="2026-07-20",
        batch="客户单-0720",
    )
    submit_shipment_report(
        db_session,
        worker.id,
        "2026-08-04",
        "源兴发",
        "裁判",
        "圆领裁判",
        [
            {"size": "M", "quantity": 10, "order_line_id": dated_order.id},
            {"size": "L", "quantity": 20, "order_line_id": batched_order.id},
        ],
        [],
        "",
    )

    response = client.get("/api/v1/shipments")

    assert response.status_code == 200
    lines = response.json()["reports"][0]["lines"]
    assert lines == [
        {
            "size": "M",
            "quantity": 10,
            "order_line_id": dated_order.id,
            "sales_order_id": None,
            "customer_sku": "",
            "system_order_no": "",
            "order_date": "2026-07-15",
            "order_batch": "",
            "order_label": "2026-07-15",
        },
        {
            "size": "L",
            "quantity": 20,
            "order_line_id": batched_order.id,
            "sales_order_id": None,
            "customer_sku": "",
            "system_order_no": "",
            "order_date": "2026-07-20",
            "order_batch": "客户单-0720",
            "order_label": "客户单-0720 · 2026-07-20",
        },
    ]


def test_shipments_photos_upload(db_session, tmp_path):
    client, _admin, worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)
    report = submit_shipment_report(
        db_session,
        worker.id,
        "2026-08-04",
        "源兴发",
        "裁判",
        "圆领裁判",
        [{"size": "M", "quantity": 10}],
        [],
        "",
    )

    response = client.post(
        f"/api/v1/shipments/{report.id}/photos",
        files={"photos": ("a.png", PNG_BYTES, "image/png")},
    )
    assert response.status_code == 200
    assert response.json()["uploaded"] == 1


def test_skus_list_update(db_session):
    client, _admin, _worker = _client(db_session)
    upsert_sku_mapping(db_session, "源兴发", "裁判", "圆领裁判", "M", "SKU-1", "6901")

    response = client.get("/api/v1/skus?q=SKU-1")
    assert response.status_code == 200
    mapping = response.json()["mappings"][0]
    assert mapping["barcode"] == "6901"

    response = client.post(
        f"/api/v1/skus/{mapping['id']}/update",
        data={
            "company_name": "源兴发",
            "product_name": "裁判",
            "style_name": "圆领裁判",
            "size": "M",
            "sku": "SKU-2",
            "barcode": "6902",
        },
    )
    assert response.status_code == 200
    response = client.get("/api/v1/skus?q=SKU-2")
    assert response.json()["mappings"][0]["sku"] == "SKU-2"


def test_daily_stats_and_goals(db_session):
    client, _admin, worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)
    report = submit_shipment_report(
        db_session,
        worker.id,
        "2026-08-04",
        "源兴发",
        "裁判",
        "圆领裁判",
        [{"size": "M", "quantity": 10}],
        [],
        "",
    )
    report.status = "auto_approved"
    db_session.commit()

    stats = client.get("/api/v1/daily-stats?ship_date=2026-08-04")
    assert stats.status_code == 200
    assert stats.json()["total"] == 10

    goals = client.get("/api/v1/goals?goal_date=2026-08-04")
    assert goals.status_code == 200
    saved = client.post("/api/v1/goals", data={"goal_date": "2026-08-04", "goal_text": "先补M码\n再清点库存"})
    assert saved.status_code == 200
    goals = client.get("/api/v1/goals?goal_date=2026-08-04")
    assert goals.json()["goal_text"] == "先补M码\n再清点库存"


def test_users_list_create_update_password(db_session):
    client, _admin, _worker = _client(db_session)

    users = client.get("/api/v1/users")
    assert users.status_code == 200
    assert len(users.json()["users"]) == 2

    created = client.post(
        "/api/v1/users",
        data={"username": "ck99", "display_name": "仓库99", "password": "pass123"},
    )
    assert created.status_code == 200
    new_user = db_session.query(User).filter_by(username="ck99").one()

    updated = client.post(
        f"/api/v1/users/{new_user.id}/update",
        data={"username": "ck99", "display_name": "仓库九九", "role": "worker", "is_active": "1"},
    )
    assert updated.status_code == 200
    assert new_user.display_name == "仓库九九"

    password = client.post(f"/api/v1/users/{new_user.id}/password", data={"password": "newpass"})
    assert password.status_code == 200


def test_logs_and_waybills(db_session, tmp_path):
    client, _admin, _worker = _client(db_session)

    logs = client.get("/api/v1/logs")
    assert logs.status_code == 200
    assert "logs" in logs.json()

    waybills = client.get("/api/v1/waybills")
    assert waybills.status_code == 200
    assert "companies" in waybills.json()

    upload = client.post(
        "/api/v1/waybills/upload",
        data={"company": "源兴发", "waybill_date": "2026-08-04"},
        files={"files": ("w.png", PNG_BYTES, "image/png")},
    )
    assert upload.status_code == 200
    assert upload.json()["imported"] == 1


def test_export_route_returns_xlsx(db_session):
    client, _admin, _worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)

    response = client.post(
        "/admin/export",
        data={"export_type": "company", "company": "源兴发", "template": "customer"},
    )

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers.get("content-type", "")


def test_unshipped_export_route_keeps_customer_price_private_and_internal_price_visible(db_session):
    client, admin, _worker = _client(db_session)
    open_line = create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100, batch="价格测试")
    complete_line = create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "L", 20, batch="已完成")
    other_company_line = create_order_line(db_session, "其他客户", "裁判", "圆领裁判", "S", 30)
    open_line.unit_price = Decimal("12.30")
    complete_line.unit_price = Decimal("20.00")
    other_company_line.unit_price = Decimal("5.00")
    report = ShipmentReport(
        user_id=admin.id, ship_date="2026-08-04", company_name="源兴发",
        product_name="裁判", style_name="圆领裁判", status="auto_approved",
    )
    db_session.add(report)
    db_session.flush()
    db_session.add(ShipmentLine(report_id=report.id, order_line_id=complete_line.id, size="L", quantity=20))
    db_session.commit()

    for template in ("customer", "internal"):
        response = client.post(
            "/admin/export",
            data={"template": template, "company": "源兴发", "export_type": "unshipped"},
        )
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["未发货明细"]
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        if template == "customer":
            assert headers == ["公司", "产品", "款式", "订单", "下单时间", "尺码", "SKU", "下单数量", "已发数量", "未发数量", "快递单号", "客户SKU"]
            assert not {"商品单价", "单价", "金额", "订单总金额"}.intersection(headers)
        assert len(rows) == 1
        assert rows[0][headers.index("公司")] == "源兴发"
        assert rows[0][headers.index("尺码")] == "M"
        assert rows[0][headers.index("未发数量")] == 100
        if template == "customer":
            assert 12.3 not in rows[0]
            assert 1230 not in rows[0]
        else:
            assert rows[0][headers.index("商品单价")] == 12.3
            assert rows[0][headers.index("金额")] == 1230

    all_customer_response = client.post(
        "/admin/export",
        data={"template": "customer", "company": "__all__", "export_type": "unshipped"},
    )
    assert all_customer_response.status_code == 200
    all_customer_sheet = load_workbook(BytesIO(all_customer_response.content)).active
    all_headers = [cell.value for cell in all_customer_sheet[1]]
    assert not {"商品单价", "单价", "金额", "订单总金额"}.intersection(all_headers)
    assert {row[all_headers.index("尺码")] for row in all_customer_sheet.iter_rows(min_row=2, values_only=True)} == {"M", "S"}

    all_internal_response = client.post(
        "/admin/export",
        data={"template": "internal", "company": "__all__", "export_type": "unshipped"},
    )
    assert all_internal_response.status_code == 200
    all_internal_sheet = load_workbook(BytesIO(all_internal_response.content)).active
    internal_headers = [cell.value for cell in all_internal_sheet[1]]
    internal_rows = list(all_internal_sheet.iter_rows(min_row=2, values_only=True))
    assert {row[internal_headers.index("尺码")] for row in internal_rows} == {"M", "S"}
    assert next(row for row in internal_rows if row[internal_headers.index("尺码")] == "M")[internal_headers.index("金额")] == 1230

    for template, company, export_type in (
        ("customer", "源兴发", None),
        ("internal", "源兴发", "company"),
        ("customer", "__all__", ""),
        ("internal", "__all__", None),
    ):
        form = {"template": template, "company": company}
        if export_type is not None:
            form["export_type"] = export_type
        response = client.post("/admin/export", data=form)
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        balance_sheet = workbook["客户发货明细" if template == "customer" else "订单发货明细"]
        headers = [cell.value for cell in balance_sheet[1]]
        rows = list(balance_sheet.iter_rows(min_row=2, values_only=True))
        expected_sizes = {"M", "L", "S"} if company == "__all__" else {"M", "L"}
        assert {row[headers.index("尺码")] for row in rows} == expected_sizes
        if template == "customer":
            assert all(
                not {"商品单价", "单价", "金额", "订单总金额"}.intersection(cell.value for cell in sheet[1])
                for sheet in workbook.worksheets
            )
            assert all(
                cell.value not in (12.3, 1230)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(min_row=2)
                for cell in row
            )
        else:
            priced_row = next(row for row in rows if row[headers.index("尺码")] == "M")
            assert priced_row[headers.index("商品单价")] == 12.3
            assert priced_row[headers.index("金额")] == 1230


def test_export_route_rejects_unsupported_template_and_type_without_creating_workbook(db_session, tmp_path, monkeypatch):
    import app.main as main_module

    client, _admin, _worker = _client(db_session)
    create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)
    monkeypatch.setattr(main_module, "EXPORT_DIR", tmp_path)

    for form in (
        {"template": "customer", "company": "源兴发", "export_type": "bogus"},
        {"template": "bogus", "company": "源兴发", "export_type": "company"},
        {"template": "bogus", "company": "__all__", "export_type": "unshipped"},
    ):
        response = client.post("/admin/export", data=form)
        assert response.status_code == 400
        assert "spreadsheetml" not in response.headers.get("content-type", "")
        assert list(tmp_path.glob("*.xlsx")) == []


def test_daily_export_route_keeps_existing_shipment_sheet_and_date_filter(db_session, tmp_path, monkeypatch):
    import app.main as main_module

    client, admin, _worker = _client(db_session)
    line = create_order_line(db_session, "源兴发", "裁判", "圆领裁判", "M", 100)
    line.unit_price = Decimal("12.30")
    for ship_date, quantity in (("2026-08-04", 3), ("2026-08-05", 4)):
        report = ShipmentReport(
            user_id=admin.id, ship_date=ship_date, company_name="源兴发",
            product_name="裁判", style_name="圆领裁判", status="auto_approved",
        )
        db_session.add(report)
        db_session.flush()
        db_session.add(ShipmentLine(report_id=report.id, order_line_id=line.id, size="M", quantity=quantity))
    db_session.commit()
    monkeypatch.setattr(main_module, "EXPORT_DIR", tmp_path)

    for template in ("customer", "internal"):
        response = client.post(
            "/admin/export",
            data={"template": template, "export_type": "daily", "ship_date": "2026-08-04"},
        )
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["发货流水"]
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        assert headers == [
            "发货日期", "上报时间", "上报人", "公司", "产品", "款式", "尺码", "数量", "状态", "异常原因", "备注",
            "订单", "下单时间", "客户订单号", "颜色", "SPU", "客户SKU",
        ]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 1
        assert rows[0][headers.index("发货日期")] == "2026-08-04"
        assert rows[0][headers.index("数量")] == 3
        assert not {"商品单价", "单价", "金额", "订单总金额"}.intersection(headers)
        assert 12.3 not in rows[0]
        assert 1230 not in rows[0]


def test_work_info_get_and_save(db_session):
    client, _admin, _worker = _client(db_session)

    page = client.get("/api/v1/work-info?company=源兴发&product=裁判&style=圆领裁判")
    assert page.status_code == 200
    assert len(page.json()["rows"]) == 4

    saved = client.post(
        "/api/v1/work-info",
        data={
            "company_name": "源兴发",
            "product_name": "裁判",
            "style_name": "圆领裁判",
            "section_key": ["accessories", "bag"],
            "section_title": ["产品配件信息", "包装袋信息"],
            "content": ["帽子*1", "透明袋"],
            "existing_photo_path": ["", ""],
            "existing_original_name": ["", ""],
            "remove_photo": ["0", "0"],
        },
    )
    assert saved.status_code == 200
    page = client.get("/api/v1/work-info?company=源兴发&product=裁判&style=圆领裁判")
    assert page.json()["rows"][0]["content"] == "帽子*1"
